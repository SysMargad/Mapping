"""Export project tracker workbooks to privacy-minimised web assets.

Checklist rows become operational flight-register records, never trajectories.
Base and GCP UTM coordinates become map points and are assigned to the matching
licence extent. Pilot names are intentionally excluded from the public assets.
"""

from __future__ import annotations

import argparse
import json
import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from export_gpkg import utm_to_wgs84


PROJECTS = {
    "Hetsuu_Hutul_Drone_Project_Tracker.xlsx": {
        "key": "hetsuu-hutul", "project": "Hetsuu hutul", "label": "Хэцүү хөтөл",
        "licence": "XV-022905", "url": "https://docs.google.com/spreadsheets/d/1XvKLs9QCgvIjHjQbXlV7thqPu0HdLIae/edit?gid=289538975",
        "sourceModified": "2026-09-17T02:17:59.510Z",
    },
    "Artsat_Drone_Project_Tracker.xlsx": {
        "key": "artsat", "project": "Artsat", "label": "Арцат",
        "licence": "XV-021395", "url": "https://docs.google.com/spreadsheets/d/1dbk0R5iHgwgPfa8gT25wN-NhtOVLjQUZ/edit?gid=991918308",
        "sourceModified": "2026-09-17T02:21:23.955Z",
    },
    "Buduunkhad_Drone_Project_Tracker.xlsx": {
        "key": "buduunkhad", "project": "Buduunkhad", "label": "Бүдүүн хад",
        "licence": "XV-023222", "url": "https://docs.google.com/spreadsheets/d/1DwydImVmcQvj7s4loVRk_4K9-VQ7szXz/edit?gid=991918308",
        "sourceModified": "2026-09-17T02:22:04.459Z",
    },
}


def clean(value):
    return str(value).strip() if value not in (None, "") else ""


def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return clean(value)


def number(value):
    if value in (None, ""):
        return None
    try:
        result = float(value)
        return int(result) if result.is_integer() else result
    except (TypeError, ValueError):
        return None


def normalise_sensor(value: str) -> str:
    upper = value.upper()
    for sensor in ("L2", "L3", "P1", "MAGARROW", "MEDUSA"):
        if sensor in upper:
            return sensor.title() if sensor in {"Magarrow", "Medusa"} else sensor
    return value or "Unknown"


def licence_extents(path: Path) -> dict[str, tuple[float, float, float, float]]:
    data = json.loads(path.read_text(encoding="utf-8"))

    def points(value):
        if isinstance(value, list) and len(value) >= 2 and all(isinstance(item, (int, float)) for item in value[:2]):
            yield value
        elif isinstance(value, list):
            for item in value:
                yield from points(item)

    by_licence = {item["licence"]: item for item in PROJECTS.values()}
    extents = {}
    for feature in data.get("features", []):
        licence = feature.get("properties", {}).get("LICENSE")
        if licence not in by_licence:
            continue
        coords = list(points((feature.get("geometry") or {}).get("coordinates")))
        extents[by_licence[licence]["key"]] = (
            min(point[0] for point in coords), min(point[1] for point in coords),
            max(point[0] for point in coords), max(point[1] for point in coords),
        )
    missing = {item["key"] for item in PROJECTS.values()} - set(extents)
    if missing:
        raise ValueError(f"Missing licence extents for: {', '.join(sorted(missing))}")
    return extents


def assigned_project(lon: float, lat: float, extents: dict, margin: float = 0.005) -> str | None:
    exact = [key for key, (west, south, east, north) in extents.items() if west <= lon <= east and south <= lat <= north]
    if exact:
        return exact[0]
    nearby = []
    for key, (west, south, east, north) in extents.items():
        dx = max(west - lon, 0, lon - east)
        dy = max(south - lat, 0, lat - north)
        if dx <= margin and dy <= margin:
            nearby.append((math.hypot(dx, dy), key))
    return min(nearby)[1] if nearby else None


def checklist_records(path: Path, config: dict) -> list[dict]:
    ws = load_workbook(path, data_only=True, read_only=True)["Checklist"]
    current_area = config["label"]
    current_date = ""
    records = []
    for row_number, raw in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        row = list(raw) + [None] * 33
        if clean(row[1]):
            current_area = clean(row[1])
        if row[2] not in (None, ""):
            current_date = iso_date(row[2])
        mission = clean(row[4])
        if not mission:
            continue
        sensor_raw = clean(row[3])
        record = {
            "id": f"{config['key']}:{row_number}",
            "projectKey": config["key"],
            "area": current_area,
            "date": current_date or None,
            "sensor": normalise_sensor(sensor_raw),
            "sensorRaw": sensor_raw,
            "mission": mission,
            "flightCount": number(row[5]),
            "startTime": number(row[6]),
            "endTime": number(row[7]),
            "altitudeM": number(row[8]),
            "speedMps": number(row[9]),
            "lidarOverlapPct": number(row[10]),
            "forwardOverlapPct": number(row[11]),
            "sideOverlapPct": number(row[12]),
            "imageCount": number(row[13]) or 0,
            "fileCopy": clean(row[14]) or None,
            "gcp": clean(row[16]) or None,
            "lidar": clean(row[17]) or None,
            "metadata": clean(row[18]) or None,
            "weather": clean(row[19]) or None,
            "djiTerra": clean(row[20]) or None,
            "metashape": clean(row[21]) or None,
            "deliveryDate": iso_date(row[22]) or None,
            "qaqc": clean(row[23]) or None,
            "note": clean(row[24]) or None,
            "link": clean(row[25]) if clean(row[25]).startswith("http") else None,
            "dataType": "flight_register_record",
            "plannedActual": "actual_record",
            "trajectoryAvailable": False,
        }
        records.append({key: value for key, value in record.items() if value is not None})
    return records


def summarise(records: list[dict], config: dict) -> dict:
    dates = sorted({item.get("date") for item in records if item.get("date")})
    sensor_groups = defaultdict(list)
    daily_groups = defaultdict(list)
    for item in records:
        sensor_groups[item["sensor"]].append(item)
        if item.get("date"):
            daily_groups[item["date"]].append(item)

    def metric(group: list[dict]) -> dict:
        altitudes = [item["altitudeM"] for item in group if item.get("altitudeM") is not None]
        speeds = [item["speedMps"] for item in group if item.get("speedMps") is not None]
        return {
            "records": len(group),
            "days": len({item.get("date") for item in group if item.get("date")}),
            "images": int(sum(item.get("imageCount", 0) for item in group)),
            "altitudeMinM": min(altitudes) if altitudes else None,
            "altitudeMaxM": max(altitudes) if altitudes else None,
            "speedMinMps": min(speeds) if speeds else None,
            "speedMaxMps": max(speeds) if speeds else None,
        }

    return {
        **config,
        "dataType": "flight_register_summary",
        "plannedActual": "actual_record",
        "trajectoryAvailable": False,
        "records": len(records),
        "flightDays": len(dates),
        "images": int(sum(item.get("imageCount", 0) for item in records)),
        "dateFrom": dates[0] if dates else None,
        "dateTo": dates[-1] if dates else None,
        "fileCopy": dict(Counter(item.get("fileCopy") or "Not recorded" for item in records)),
        "qaqc": dict(Counter(item.get("qaqc") or "Not recorded" for item in records)),
        "processing": {
            "djiTerra": dict(Counter(item.get("djiTerra") or "Not recorded" for item in records)),
            "metashape": dict(Counter(item.get("metashape") or "Not recorded" for item in records)),
        },
        "areas": dict(Counter(item.get("area") or config["label"] for item in records)),
        "sensors": {key: metric(group) for key, group in sorted(sensor_groups.items())},
        "daily": [
            {
                "date": key,
                "records": len(group),
                "images": int(sum(item.get("imageCount", 0) for item in group)),
                "sensors": dict(Counter(item["sensor"] for item in group)),
            }
            for key, group in sorted(daily_groups.items(), reverse=True)
        ],
    }


def coordinate_observations(paths: list[Path], extents: dict) -> tuple[list[dict], Counter]:
    observations = []
    dropped = Counter()
    layouts = {
        "base": {"sheet": "Base", "zone": 5, "north": 7, "east": 8, "height": 9, "name": 3, "date": 1},
        "gcp": {"sheet": "GCP", "zone": 4, "north": 6, "east": 7, "height": 8, "name": 3, "date": 1},
    }
    for path in paths:
        source_config = PROJECTS[path.name]
        workbook = load_workbook(path, data_only=True, read_only=True)
        for control_type, layout in layouts.items():
            current_date = ""
            for row_number, raw in enumerate(workbook[layout["sheet"]].iter_rows(min_row=4, values_only=True), start=4):
                row = list(raw) + [None] * 31
                if row[layout["date"]] not in (None, ""):
                    current_date = iso_date(row[layout["date"]])
                zone_match = re.search(r"(4[67])", clean(row[layout["zone"]]))
                try:
                    northing = float(row[layout["north"]])
                    easting = float(row[layout["east"]])
                except (TypeError, ValueError):
                    continue
                if not zone_match:
                    dropped["missing_zone"] += 1
                    continue
                zone = int(zone_match.group(1))
                lon, lat = utm_to_wgs84(easting, northing, zone, True)
                project_key = assigned_project(lon, lat, extents)
                if project_key is None:
                    dropped["outside_project_licences"] += 1
                    continue
                observations.append({
                    "projectKey": project_key,
                    "controlType": control_type,
                    "name": clean(row[layout["name"]]) or f"{control_type.upper()} {row_number}",
                    "date": current_date or None,
                    "lon": lon, "lat": lat,
                    "northing": northing, "easting": easting,
                    "elevationM": number(row[layout["height"]]),
                    "zone": zone,
                    "sourceFile": path.name,
                    "sourceUrl": source_config["url"],
                    "sourceMatchesProject": source_config["key"] == project_key,
                })

    grouped = defaultdict(list)
    for item in observations:
        key = (
            item["projectKey"], item["controlType"], item["name"].lower(),
            round(item["lon"], 6), round(item["lat"], 6),
        )
        grouped[key].append(item)

    features = []
    for index, group in enumerate(grouped.values(), start=1):
        preferred = next((item for item in group if item["sourceMatchesProject"]), group[0])
        project = next(item for item in PROJECTS.values() if item["key"] == preferred["projectKey"])
        dates = sorted({item["date"] for item in group if item.get("date")})
        props = {
            "project": project["project"],
            "projectKey": project["key"],
            "area": project["label"],
            "licence": project["licence"],
            "sensor": "Base / Control",
            "dataType": "control_reference_point",
            "controlType": preferred["controlType"],
            "name": preferred["name"],
            "plannedActual": "reference",
            "status": "confirmed_source",
            "contextOnly": True,
            "sourceFile": preferred["sourceFile"],
            "sourceUrl": preferred["sourceUrl"],
            "sourceCrs": f"EPSG:326{preferred['zone']}",
            "displayCrs": "EPSG:4326",
            "crsVerified": True,
            "easting": preferred["easting"],
            "northing": preferred["northing"],
            "elevationM": preferred["elevationM"],
            "dates": dates,
            "observationCount": len(group),
        }
        features.append({
            "type": "Feature", "id": f"tracker-control:{index}",
            "properties": {key: value for key, value in props.items() if value not in (None, [], "")},
            "geometry": {"type": "Point", "coordinates": [preferred["lon"], preferred["lat"]]},
        })
    return features, dropped


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbooks", nargs=3, type=Path)
    parser.add_argument("--licences", required=True, type=Path)
    parser.add_argument("--summary-output", required=True, type=Path)
    parser.add_argument("--points-output", required=True, type=Path)
    args = parser.parse_args()
    paths = {path.name: path for path in args.workbooks}
    missing = set(PROJECTS) - set(paths)
    if missing:
        raise ValueError(f"Missing source workbook(s): {', '.join(sorted(missing))}")
    ordered_paths = [paths[name] for name in PROJECTS]
    extents = licence_extents(args.licences)

    all_records = []
    projects = {}
    for path in ordered_paths:
        config = PROJECTS[path.name]
        records = checklist_records(path, config)
        projects[config["key"]] = summarise(records, config)
        all_records.extend(records)

    features, dropped = coordinate_observations(ordered_paths, extents)
    point_counts = defaultdict(Counter)
    for feature in features:
        props = feature["properties"]
        point_counts[props["projectKey"]][props["controlType"]] += 1
    for key, project in projects.items():
        project["controlPoints"] = dict(point_counts[key])

    summary = {
        "project": "Multi-project operations",
        "scope": "project_operations",
        "dataType": "flight_register_summary",
        "generated": datetime.now().astimezone().isoformat(timespec="seconds"),
        "privacy": "Pilot names are excluded from this web asset.",
        "trajectoryNote": "Checklist records confirm flights but do not contain flown trajectory geometry.",
        "projects": projects,
        "records": all_records,
    }
    points = {
        "type": "FeatureCollection",
        "project": "Multi-project operations",
        "scope": "project_operations",
        "dataType": "control_reference_point",
        "featureCount": len(features),
        "dropped": dict(dropped),
        "features": features,
    }
    args.summary_output.parent.mkdir(parents=True, exist_ok=True)
    args.points_output.parent.mkdir(parents=True, exist_ok=True)
    args.summary_output.write_text(json.dumps(summary, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    args.points_output.write_text(json.dumps(points, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Exported {len(all_records)} flight records and {len(features)} unique control points")
    for key, project in projects.items():
        print(f"- {key}: {project['records']} records, {project['flightDays']} days, {project['images']} images, {project['controlPoints']}")
    print(f"Dropped coordinate observations: {dict(dropped)}")


if __name__ == "__main__":
    main()
