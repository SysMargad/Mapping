"""Synchronise verified project flight sources from Google Drive.

This script is intended for GitHub Actions with a read-only Google Drive
service account. It downloads the three tracker workbooks, follows each
Checklist mission folder hyperlink, imports only coordinate-bearing flight
files, rebuilds coverage, and leaves git commit/push to the workflow.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import coordinate_sources as cs  # noqa: E402


FOLDER_MIME = "application/vnd.google-apps.folder"
SHORTCUT_MIME = "application/vnd.google-apps.shortcut"
SHEET_MIME = "application/vnd.google-apps.spreadsheet"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SUPPORTED_SUFFIXES = {".mrk", ".kmz", ".kml", ".csv", ".txt"}
SOURCE_PRIORITY = {".kmz": 0, ".mrk": 1, ".kml": 2, ".csv": 3, ".txt": 4}
FOLDER_ID_RE = re.compile(r"/folders/([A-Za-z0-9_-]+)")
DATE_COMPACT_RE = re.compile(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)")
DATE_DASH_RE = re.compile(r"(?<!\d)(20\d{2})-(\d{2})-(\d{2})(?!\d)")
DJI_FILE_RE = re.compile(r"^DJI(?:[_\-\s]|$)", re.IGNORECASE)
RAW_DATA_FOLDER_KEYS = {"0rawdata", "rawdata"}
DJI_FILE_COUNT_SENSORS = {4: "P1", 11: "L3"}
DRONE_FOLDER_RE = re.compile(
    r"(?:^|[^a-z0-9])(drone|uav|magnetic[ _-]*survey|magarrow|flight[ _-]*(?:log|record))(?:$|[^a-z0-9])",
    re.IGNORECASE,
)
DRONE_SOURCE_RE = re.compile(
    r"(?:DJIFlightRecord|Timestamp|(?:^|[_-])DJI(?:[_-]|$)|SRVY\d*[-_]ACQU\d*|flight[ _-]*(?:log|record|track)|trajectory)",
    re.IGNORECASE,
)
# Name patterns alone cannot decide whether a CSV/TXT holds coordinates, so
# unrecognised tables inside a configured survey branch are probed by header.
PROBE_SUFFIXES = {".csv", ".txt"}
PROBE_BYTES = 64 * 1024
MAX_HEADER_PROBES = 400


def clean(value) -> str:
    return str(value).strip() if value not in (None, "") else ""


def iso_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    match = re.match(r"(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})", text)
    return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}" if match else None


def source_date(name: str) -> str | None:
    match = DATE_DASH_RE.search(name) or DATE_COMPACT_RE.search(name)
    if not match:
        return None
    return f"{match.group(1)}-{match.group(2)}-{match.group(3)}"


def folder_id(value: str | None) -> str | None:
    match = FOLDER_ID_RE.search(value or "")
    return match.group(1) if match else None


def safe_filename(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._") or "source"


def comparable_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def tracker_rows(path: Path, project_key: str) -> list[dict]:
    # Imported here so the discovery and matching helpers stay importable (and
    # testable) on a runner that has no Drive/Excel dependencies installed.
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet = workbook["Checklist"]
    current_date = None
    rows = []
    for row_number in range(3, sheet.max_row + 1):
        current_date = iso_date(sheet.cell(row_number, 3).value) or current_date
        mission = clean(sheet.cell(row_number, 5).value)
        if not mission:
            continue
        link_cell = sheet.cell(row_number, 26)
        link = link_cell.hyperlink.target if link_cell.hyperlink else clean(link_cell.value)
        linked_folder = folder_id(link)
        if not linked_folder:
            continue
        rows.append({
            "trackerId": f"{project_key}:{row_number}",
            "mission": mission,
            "date": current_date,
            "folderId": linked_folder,
        })
    return rows


def build_drive_service():
    raw = os.environ.get("GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON", "").strip()
    if not raw:
        raise RuntimeError("GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON is not configured")
    try:
        service_account_info = json.loads(raw)
    except json.JSONDecodeError as error:
        raise RuntimeError("GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON is not valid JSON") from error
    from google.oauth2 import service_account
    from googleapiclient.discovery import build

    credentials = service_account.Credentials.from_service_account_info(
        service_account_info,
        scopes=["https://www.googleapis.com/auth/drive.readonly"],
    )
    return build("drive", "v3", credentials=credentials, cache_discovery=False)


def metadata(service, file_id: str) -> dict:
    return service.files().get(
        fileId=file_id,
        fields="id,name,mimeType,size,modifiedTime,md5Checksum,shortcutDetails,webViewLink",
        supportsAllDrives=True,
    ).execute()


def list_children(service, parent_id: str) -> list[dict]:
    result = []
    token = None
    while True:
        response = service.files().list(
            q=f"'{parent_id}' in parents and trashed = false",
            fields="nextPageToken,files(id,name,mimeType,size,modifiedTime,md5Checksum,shortcutDetails,webViewLink)",
            pageSize=1000,
            pageToken=token,
            spaces="drive",
            includeItemsFromAllDrives=True,
            supportsAllDrives=True,
        ).execute()
        result.extend(response.get("files", []))
        token = response.get("nextPageToken")
        if not token:
            return result


def resolve_shortcut(service, item: dict) -> dict:
    if item.get("mimeType") != SHORTCUT_MIME:
        return item
    target_id = item.get("shortcutDetails", {}).get("targetId")
    return metadata(service, target_id) if target_id else item


def find_sources(service, root_folder_id: str, max_depth: int) -> list[dict]:
    found = []
    queue = [(root_folder_id, 0)]
    visited = set()
    while queue:
        current, depth = queue.pop(0)
        if current in visited:
            continue
        visited.add(current)
        for raw_item in list_children(service, current):
            item = resolve_shortcut(service, raw_item)
            if item.get("mimeType") == FOLDER_MIME:
                if depth < max_depth:
                    queue.append((item["id"], depth + 1))
                continue
            if Path(item.get("name", "")).suffix.lower() in SUPPORTED_SUFFIXES:
                found.append(item)
    return found


def source_date_from_path(value: str) -> str | None:
    """Read a full date, or a YYYY ancestor plus MMDD child folder, from a Drive path."""
    direct = source_date(value)
    if direct:
        return direct
    parts = [part for part in re.split(r"[/\\]+", value) if part]
    for index, part in enumerate(parts):
        year_match = re.fullmatch(r"(?:\d+[._ -]*)?(20\d{2})", part)
        if not year_match:
            continue
        year = year_match.group(1)
        for child in parts[index + 1:]:
            day_match = re.match(r"(\d{2})(\d{2})(?:\D|$)", child)
            if not day_match:
                continue
            try:
                return date(int(year), int(day_match.group(1)), int(day_match.group(2))).isoformat()
            except ValueError:
                continue
    return None


def is_named_drone_source(name: str) -> bool:
    """Fast path: the file name itself identifies a known flight-source format."""
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        return False
    if suffix == ".mrk":
        return True
    return bool(DRONE_SOURCE_RE.search(Path(name).name))


def probe_header_text(service, item: dict) -> str | None:
    """Download the first PROBE_BYTES of a file so its header can be read.

    Returns None when the bytes cannot be read or decoded. The file is only
    read; nothing in Drive is modified.
    """
    from googleapiclient.http import MediaIoBaseDownload

    try:
        request = service.files().get_media(fileId=item["id"], supportsAllDrives=True)
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request, chunksize=PROBE_BYTES)
        downloader.next_chunk()
        return buffer.getvalue().decode("utf-8-sig", errors="replace")
    except Exception:
        return None


def classify_candidate(service, item: dict, probe_budget: list[int]) -> tuple[bool, str]:
    """Decide whether a Drive item is a coordinate source. Returns (keep, how)."""
    name = item.get("name", "")
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        return False, "unsupported-suffix"
    if is_named_drone_source(name):
        return True, "name-pattern"
    if suffix not in PROBE_SUFFIXES:
        return False, "name-pattern-miss"
    if probe_budget[0] <= 0:
        return False, "probe-budget-exhausted"
    probe_budget[0] -= 1
    text = probe_header_text(service, item)
    if text is None:
        return False, "probe-unreadable"
    header = cs.read_header(text)
    if header is None:
        return False, "probe-no-header"
    if cs.looks_like_point_table(header):
        # A GCP/sample coordinate table is not a flown trajectory.
        return False, "probe-point-table"
    if cs.detect_columns(header) is None:
        return False, "probe-no-coordinates"
    return True, "probe-header"


def find_drone_source_roots(service, root_folder_id: str, max_depth: int = 5) -> list[dict]:
    """Locate explicitly named drone/UAV/magnetic-survey branches under a project root."""
    found = []
    queue = [(root_folder_id, "", 0)]
    visited = set()
    while queue:
        current, parent_path, depth = queue.pop(0)
        if current in visited:
            continue
        visited.add(current)
        for raw_item in list_children(service, current):
            item = resolve_shortcut(service, raw_item)
            if item.get("mimeType") != FOLDER_MIME:
                continue
            path = f"{parent_path}/{item.get('name', '')}".strip("/")
            if DRONE_FOLDER_RE.search(item.get("name", "")):
                found.append({**item, "relativePath": path})
                # The matching branch is scanned separately below, so do not
                # rediscover every nested raw/processed folder as another root.
                continue
            if depth < max_depth:
                queue.append((item["id"], path, depth + 1))
    return found


def find_drone_sources(service, root_folder_id: str, discovery_depth: int = 5, scan_depth: int = 8) -> tuple[list[dict], dict]:
    """Return strong coordinate-source candidates without exposing private paths in the summary."""
    roots = find_drone_source_roots(service, root_folder_id, discovery_depth)
    found = []
    seen_files = set()
    scanned_folders = 0
    inspected_files = 0
    probe_budget = [MAX_HEADER_PROBES]
    discovery_reasons: dict[str, int] = {}
    for root in roots:
        queue = [(root["id"], root.get("relativePath", root.get("name", "")), 0)]
        visited = set()
        while queue:
            current, parent_path, depth = queue.pop(0)
            if current in visited:
                continue
            visited.add(current)
            scanned_folders += 1
            for raw_item in list_children(service, current):
                item = resolve_shortcut(service, raw_item)
                path = f"{parent_path}/{item.get('name', '')}".strip("/")
                if item.get("mimeType") == FOLDER_MIME:
                    if depth < scan_depth:
                        queue.append((item["id"], path, depth + 1))
                    continue
                if item.get("id") in seen_files:
                    continue
                inspected_files += 1
                keep, how = classify_candidate(service, item, probe_budget)
                discovery_reasons[how] = discovery_reasons.get(how, 0) + 1
                if not keep:
                    continue
                seen_files.add(item["id"])
                found.append({
                    **item,
                    "relativePath": path,
                    "sourceDate": source_date_from_path(path),
                    "discoveredBy": how,
                })
    dated = sorted(item["sourceDate"] for item in found if item.get("sourceDate"))
    suffix_counts = {}
    for item in found:
        suffix = Path(item.get("name", "")).suffix.lower().lstrip(".") or "unknown"
        suffix_counts[suffix] = suffix_counts.get(suffix, 0) + 1
    summary = {
        "matchedBranchCount": len(roots),
        "scannedFolderCount": scanned_folders,
        "inspectedFileCount": inspected_files,
        "headerProbesUsed": MAX_HEADER_PROBES - probe_budget[0],
        "coordinateSourceCount": len(found),
        "sourceCountByType": dict(sorted(suffix_counts.items())),
        "discoveryDecisions": dict(sorted(discovery_reasons.items())),
        "dateFrom": dated[0] if dated else None,
        "dateTo": dated[-1] if dated else None,
    }
    return found, summary


def match_campaign(campaigns: list[dict], project_key: str, candidate: dict) -> dict | None:
    """Return the one campaign that explicitly claims this source, else None.

    A campaign must declare the source; nothing is adopted because of where it
    happens to sit. Two competing campaigns resolve to no match.
    """
    name = candidate.get("name", "")
    path = candidate.get("relativePath", name)
    matched = []
    for campaign in campaigns:
        if campaign.get("projectKey") != project_key:
            continue
        rules = campaign.get("match", {})
        folders = rules.get("pathContains", [])
        if folders and not any(token.lower() in path.lower() for token in folders):
            continue
        patterns = rules.get("fileNamePatterns", [])
        if patterns and not any(re.search(pattern, name, re.IGNORECASE) for pattern in patterns):
            continue
        if not folders and not patterns:
            continue
        matched.append(campaign)
    return matched[0] if len(matched) == 1 else None


def folder_name_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def find_raw_data_folder(service, root_folder_id: str, max_depth: int = 2) -> dict | None:
    queue = [(root_folder_id, 0)]
    visited = set()
    while queue:
        current, depth = queue.pop(0)
        if current in visited:
            continue
        visited.add(current)
        for raw_item in list_children(service, current):
            item = resolve_shortcut(service, raw_item)
            if item.get("mimeType") != FOLDER_MIME:
                continue
            if folder_name_key(item.get("name", "")) in RAW_DATA_FOLDER_KEYS:
                return item
            if depth < max_depth:
                queue.append((item["id"], depth + 1))
    return None


def raw_sensor_counts(service, root_folder_id: str, max_depth: int = 8) -> dict:
    """Return aggregate sensor counts without exposing private folder metadata."""
    raw_folder = find_raw_data_folder(service, root_folder_id)
    if not raw_folder:
        raise ValueError("Raw Data folder was not found under the configured root")
    sensors = {
        sensor: {"folderCount": 0, "djiFileCount": 0}
        for sensor in sorted(set(DJI_FILE_COUNT_SENSORS.values()))
    }
    unclassified_counts = {}
    scanned_folders = 0
    queue = [(raw_folder["id"], 0)]
    visited = set()
    while queue:
        current, depth = queue.pop(0)
        if current in visited:
            continue
        visited.add(current)
        scanned_folders += 1
        children = [resolve_shortcut(service, item) for item in list_children(service, current)]
        child_folders = [item for item in children if item.get("mimeType") == FOLDER_MIME]
        dji_file_count = sum(
            1 for item in children
            if item.get("mimeType") != FOLDER_MIME and DJI_FILE_RE.search(item.get("name", ""))
        )
        if dji_file_count:
            sensor = DJI_FILE_COUNT_SENSORS.get(dji_file_count)
            if sensor:
                sensors[sensor]["folderCount"] += 1
                sensors[sensor]["djiFileCount"] += dji_file_count
            else:
                key = str(dji_file_count)
                unclassified_counts[key] = unclassified_counts.get(key, 0) + 1
        if depth < max_depth:
            queue.extend((item["id"], depth + 1) for item in child_folders)
    return {
        "classificationRule": {str(count): sensor for count, sensor in DJI_FILE_COUNT_SENSORS.items()},
        "scannedFolderCount": scanned_folders,
        "sensors": sensors,
        "unclassifiedFolderCountsByDjiFileCount": unclassified_counts,
    }


def choose_source(
    items: list[dict],
    tracker_date: str | None,
    mission: str,
    max_size: int,
    used_source_ids: set[str],
) -> dict | None:
    eligible = []
    for item in items:
        suffix = Path(item.get("name", "")).suffix.lower()
        item_date = source_date(item.get("name", ""))
        try:
            size = int(item.get("size") or 0)
        except (TypeError, ValueError):
            size = 0
        if item.get("id") in used_source_ids or suffix not in SUPPORTED_SUFFIXES or size > max_size:
            continue
        if tracker_date and item_date and item_date != tracker_date:
            continue
        eligible.append(item)
    if not eligible:
        return None

    # A linked folder can contain more than one mission from the same day. Prefer
    # the file whose name contains the complete tracker mission token, when that
    # is available, before applying the format preference.
    mission_key = comparable_name(mission)
    named_matches = [item for item in eligible if mission_key and mission_key in comparable_name(item.get("name", ""))]
    if named_matches:
        eligible = named_matches

    # Prefer a newly modified file within the best supported source type.
    eligible.sort(key=lambda item: item.get("modifiedTime", ""), reverse=True)
    eligible.sort(key=lambda item: SOURCE_PRIORITY.get(Path(item.get("name", "")).suffix.lower(), 99))
    return eligible[0]


def download_file(service, item: dict, target: Path) -> None:
    from googleapiclient.http import MediaIoBaseDownload

    target.parent.mkdir(parents=True, exist_ok=True)
    if item.get("mimeType") == SHEET_MIME:
        request = service.files().export_media(fileId=item["id"], mimeType=XLSX_MIME)
    else:
        request = service.files().get_media(fileId=item["id"], supportsAllDrives=True)
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request, chunksize=4 * 1024 * 1024)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    target.write_bytes(buffer.getvalue())


def run(command: list[str], cwd: Path) -> None:
    print("Running:", " ".join(command))
    subprocess.run(command, cwd=cwd, check=True)


def sync(config_path: Path, repo: Path, workspace: Path) -> dict:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    projects = config.get("projects", [])
    required_keys = {"hetsuu-hutul", "artsat", "buduunkhad"}
    if {item.get("key") for item in projects} != required_keys:
        raise ValueError("drive-sync-sources.json must configure exactly the three project trackers")
    if workspace.exists():
        shutil.rmtree(workspace)
    workbooks_dir = workspace / "workbooks"
    sources_dir = workspace / "sources"
    reports_dir = workspace / "reports"
    workbooks_dir.mkdir(parents=True)
    sources_dir.mkdir(parents=True)
    reports_dir.mkdir(parents=True)
    service = build_drive_service()
    raw_root_folder_id = os.environ.get("NERGUI_UNDUR_RAW_ROOT_FOLDER_ID", "").strip()
    hetsuu_root_folder_id = os.environ.get("HETSUU_HUTUL_ROOT_FOLDER_ID", "").strip()
    raw_sensor_summary = None
    if raw_root_folder_id:
        try:
            raw_sensor_summary = raw_sensor_counts(service, raw_root_folder_id)
            print("Nergui Undur Raw Data sensor classification:")
            print(json.dumps(raw_sensor_summary, ensure_ascii=False, indent=2))
        except Exception as error:
            print(f"::warning::Nergui Undur Raw Data sensor classification skipped: {error}")
    else:
        print("::notice::NERGUI_UNDUR_RAW_ROOT_FOLDER_ID is not configured; Raw Data sensor classification skipped.")
    hetsuu_root_sources = []
    hetsuu_drone_summary = None
    if hetsuu_root_folder_id:
        try:
            hetsuu_root_sources, hetsuu_drone_summary = find_drone_sources(
                service,
                hetsuu_root_folder_id,
                int(config.get("projectRootDiscoveryDepth", 5)),
                int(config.get("projectRootScanDepth", 8)),
            )
            print("Hetsuu Hutul project-root drone source discovery:")
            print(json.dumps(hetsuu_drone_summary, ensure_ascii=False, indent=2))
        except Exception as error:
            print(f"::warning::Hetsuu Hutul project-root drone source discovery skipped: {error}")
    else:
        print("::notice::HETSUU_HUTUL_ROOT_FOLDER_ID is not configured; project-root drone discovery skipped.")
    max_depth = int(config.get("folderScanDepth", 2))
    max_size = int(config.get("maxSourceSizeBytes", 100 * 1024 * 1024))
    campaigns = config.get("campaigns", [])
    workbook_paths = []
    project_results = {}

    for project in projects:
        tracker_item = metadata(service, project["trackerFileId"])
        workbook_path = workbooks_dir / project["workbookName"]
        download_file(service, tracker_item, workbook_path)
        workbook_paths.append(workbook_path)
        rows = tracker_rows(workbook_path, project["key"])
        project_source_dir = sources_dir / project["key"]
        project_source_dir.mkdir(parents=True)
        mapping_entries = []
        seen_folders = {}
        used_source_ids: set[str] = set()
        mismatched = []
        stage = {
            "missionFoldersScanned": 0,
            "missionFoldersUnreadable": 0,
            "candidateFilesFound": 0,
            "sourcesChosen": 0,
            "sourcesDownloaded": 0,
            "downloadFailures": 0,
        }
        reasons: dict[str, int] = {}

        def note(code: str) -> None:
            reasons[code] = reasons.get(code, 0) + 1

        for row in rows:
            candidates = seen_folders.get(row["folderId"])
            if candidates is None:
                try:
                    candidates = find_sources(service, row["folderId"], max_depth)
                    stage["missionFoldersScanned"] += 1
                except Exception as error:
                    # A folder the service account cannot read is an access
                    # problem, not a parsing bug.
                    candidates = []
                    stage["missionFoldersUnreadable"] += 1
                    print(f"::warning::{project['key']} mission folder unreadable ({type(error).__name__})")
                seen_folders[row["folderId"]] = candidates
                stage["candidateFilesFound"] += len(candidates)
            chosen = choose_source(candidates, row["date"], row["mission"], max_size, used_source_ids)
            if not chosen:
                wrong_dates = sorted({source_date(item.get("name", "")) for item in candidates if source_date(item.get("name", ""))})
                if wrong_dates:
                    mismatched.append({"trackerId": row["trackerId"], "date": row["date"], "sourceDates": wrong_dates})
                    note(cs.DATE_MISMATCH)
                elif not candidates:
                    note(cs.MISSING_SOURCE)
                else:
                    note(cs.MISSION_UNMATCHED)
                continue
            stage["sourcesChosen"] += 1
            used_source_ids.add(chosen["id"])
            local_name = f"{row['trackerId'].split(':')[-1].zfill(4)}__{chosen['id']}__{safe_filename(chosen['name'])}"
            local_path = project_source_dir / local_name
            try:
                download_file(service, chosen, local_path)
            except Exception as error:
                stage["downloadFailures"] += 1
                note(cs.ACCESS_DENIED)
                print(f"::warning::{project['key']} source download failed ({type(error).__name__})")
                continue
            stage["sourcesDownloaded"] += 1
            mapping_entries.append({
                "file": local_name,
                "trackerId": row["trackerId"],
                "mission": row["mission"],
                "sourceUrl": f"https://drive.google.com/file/d/{chosen['id']}/view",
            })
        campaign_entries: list[dict] = []
        if project["key"] == "hetsuu-hutul" and hetsuu_root_sources:
            tracker_missions = {
                comparable_name(row["mission"]): row
                for row in rows
                if comparable_name(row.get("mission", ""))
            }
            campaign_source_dir = sources_dir / f"{project['key']}-campaign"
            campaign_source_dir.mkdir(parents=True, exist_ok=True)
            for candidate in hetsuu_root_sources:
                if candidate.get("id") in used_source_ids:
                    continue
                candidate_key = comparable_name(candidate.get("relativePath", candidate.get("name", "")))
                exact_rows = [row for key, row in tracker_missions.items() if key in candidate_key]
                try:
                    size = int(candidate.get("size") or 0)
                except (TypeError, ValueError):
                    size = 0
                if size > max_size:
                    continue
                if len(exact_rows) > 1:
                    note(cs.AMBIGUOUS_MATCH)
                    continue
                campaign = match_campaign(campaigns, project["key"], candidate)
                if not exact_rows and not campaign:
                    # Not a 2026 tracker mission and not declared by any
                    # campaign, so it stays an inventory entry only. A source is
                    # never imported on the strength of its location alone.
                    note(cs.MISSION_UNMATCHED)
                    continue
                used_source_ids.add(candidate["id"])
                if exact_rows:
                    local_name = f"root__{candidate['id']}__{safe_filename(candidate['name'])}"
                    local_path = project_source_dir / local_name
                else:
                    local_name = f"{campaign['id']}__{candidate['id']}__{safe_filename(candidate['name'])}"
                    local_path = campaign_source_dir / local_name
                try:
                    download_file(service, candidate, local_path)
                except Exception as error:
                    stage["downloadFailures"] += 1
                    note(cs.ACCESS_DENIED)
                    print(f"::warning::{project['key']} root source download failed ({type(error).__name__})")
                    continue
                stage["sourcesDownloaded"] += 1
                entry = {
                    "file": local_name,
                    "sourceUrl": f"https://drive.google.com/file/d/{candidate['id']}/view",
                    "driveFileId": candidate["id"],
                    "md5Checksum": candidate.get("md5Checksum"),
                    "sizeBytes": size,
                    "modifiedTime": candidate.get("modifiedTime"),
                    "discoveredBy": candidate.get("discoveredBy"),
                }
                if exact_rows:
                    entry.update({
                        "trackerId": exact_rows[0]["trackerId"],
                        "mission": exact_rows[0]["mission"],
                    })
                    mapping_entries.append(entry)
                else:
                    entry["campaignId"] = campaign["id"]
                    entry["folderDate"] = candidate.get("sourceDate")
                    campaign_entries.append(entry)
        mapping_path = workspace / f"{project['key']}-mapping.json"
        mapping_path.write_text(json.dumps({"sources": mapping_entries}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        project_results[project["key"]] = {
            "trackerMissionRows": len(rows),
            "trackerRowsWithFolders": len(rows),
            "stages": stage,
            "reasons": dict(sorted(reasons.items())),
            "coordinateSources": len(mapping_entries),
            "campaignSources": len(campaign_entries),
            "dateMismatches": mismatched,
            "sourceDir": project_source_dir,
            "mapping": mapping_path,
        }
        if campaign_entries:
            campaign_mapping_path = workspace / f"{project['key']}-campaign-mapping.json"
            campaign_mapping_path.write_text(
                json.dumps({"sources": campaign_entries}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            project_results[project["key"]]["campaignSourceDir"] = sources_dir / f"{project['key']}-campaign"
            project_results[project["key"]]["campaignMapping"] = campaign_mapping_path
        if project["key"] == "hetsuu-hutul" and hetsuu_drone_summary is not None:
            project_results[project["key"]]["projectRootDroneDiscovery"] = hetsuu_drone_summary

    context_dir = repo / "dist" / "data" / "context"
    python = sys.executable
    run([
        python, str(repo / "tools" / "export_project_trackers.py"),
        *[str(path) for path in workbook_paths],
        "--licences", str(context_dir / "licenses.geojson"),
        "--summary-output", str(context_dir / "project-trackers.json"),
        "--points-output", str(context_dir / "project-control-points.geojson"),
    ], repo)

    tracks = context_dir / "project-flight-tracks.geojson"
    for project in projects:
        result = project_results[project["key"]]
        report = reports_dir / f"{project['key']}.json"
        run([
            python, str(repo / "tools" / "import_project_flight_tracks.py"),
            str(result["sourceDir"]),
            "--project", project["key"],
            "--trackers", str(context_dir / "project-trackers.json"),
            "--licences", str(context_dir / "licenses.geojson"),
            "--existing", str(tracks),
            "--output", str(tracks),
            "--mapping", str(result["mapping"]),
            "--report", str(report),
        ], repo)
        import_report = json.loads(report.read_text(encoding="utf-8"))
        result["importReport"] = {key: value for key, value in import_report.items() if key != "audit"}
        # The per-file audit names Drive files and folders, so it stays in the
        # private workflow log and never reaches a published asset.
        print(f"{project['key']} tracker import audit:")
        print(json.dumps(import_report.get("audit", []), ensure_ascii=False, indent=2))
        result.pop("sourceDir")
        result.pop("mapping")

    campaign_tracks = context_dir / "project-campaign-tracks.geojson"
    for project in projects:
        result = project_results[project["key"]]
        campaign_dir = result.pop("campaignSourceDir", None)
        campaign_mapping = result.pop("campaignMapping", None)
        if not campaign_dir or not campaign_mapping:
            continue
        campaign_report = reports_dir / f"{project['key']}-campaign.json"
        run([
            python, str(repo / "tools" / "import_campaign_tracks.py"),
            str(campaign_dir),
            "--project", project["key"],
            "--campaigns", str(config_path),
            "--licences", str(context_dir / "licenses.geojson"),
            "--mapping", str(campaign_mapping),
            "--existing", str(campaign_tracks),
            "--output", str(campaign_tracks),
            "--audit", str(campaign_report),
        ], repo)
        payload = json.loads(campaign_report.read_text(encoding="utf-8"))
        result["campaignImportReport"] = {key: value for key, value in payload.items() if key != "audit"}
        print(f"{project['key']} campaign import audit:")
        print(json.dumps(payload.get("audit", []), ensure_ascii=False, indent=2))

    run([
        python, str(repo / "tools" / "export_project_flight_coverage.py"),
        str(tracks),
        "--licences", str(context_dir / "licenses.geojson"),
        "--output", str(context_dir / "project-flight-coverage.json"),
    ], repo)
    summary = {
        "projects": project_results,
        "nerguiUndurRawSensorCounts": raw_sensor_summary,
        "hetsuuHutulDroneSources": hetsuu_drone_summary,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", required=True, type=Path)
    parser.add_argument("--repo", required=True, type=Path)
    parser.add_argument("--workspace", required=True, type=Path)
    args = parser.parse_args()
    sync(args.config.resolve(), args.repo.resolve(), args.workspace.resolve())


if __name__ == "__main__":
    main()
